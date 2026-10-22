"""
test_splitter.py
================
Unit tests for WordSplitter core functionality.
"""

import os
import shutil
import unittest
import docx
from docx.shared import Pt, RGBColor
from word_splitter_core import WordSplitter


class TestWordSplitter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "test_temp")
        os.makedirs(cls.test_dir, exist_ok=True)

        cls.sample_path = os.path.join(cls.test_dir, "sample_document.docx")
        cls.create_sample_doc(cls.sample_path)

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(cls.test_dir):
            shutil.rmtree(cls.test_dir, ignore_errors=True)

    @classmethod
    def create_sample_doc(cls, path):
        doc = docx.Document()
        
        # Section 1
        h1 = doc.add_heading("Chapter 1: The Beginning", level=1)
        p1 = doc.add_paragraph("This is the first paragraph of chapter one. It contains several important words.")
        r1 = p1.add_run(" Here is bold and italic text.")
        r1.bold = True
        r1.italic = True
        r1.font.color.rgb = RGBColor(0, 102, 204)

        p2 = doc.add_paragraph("Second paragraph continues the introduction with more descriptive details.")

        # Table
        tbl = doc.add_table(rows=2, cols=3)
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.text = f"R{r_idx}C{c_idx}"

        # Section 2
        h2 = doc.add_heading("Chapter 2: The Core Analysis", level=1)
        p3 = doc.add_paragraph("Chapter two discusses the methodology, experimental design, and data collection.")
        p4 = doc.add_paragraph("Additional explanatory notes regarding the experiments performed.")

        # Section 3
        h3 = doc.add_heading("Chapter 3: Final Conclusions", level=1)
        p5 = doc.add_paragraph("In conclusion, the document splitter separates files while maintaining style integrity.")
        p6 = doc.add_paragraph("Final remarks and closing acknowledgements.")

        doc.save(path)

    def test_analysis(self):
        analysis = WordSplitter.analyze_document(self.sample_path)
        self.assertEqual(analysis["file_name"], "sample_document.docx")
        self.assertGreater(analysis["total_paragraphs"], 5)
        self.assertEqual(analysis["total_tables"], 1)
        self.assertEqual(len(analysis["headings"]), 3)
        self.assertGreater(analysis["total_words"], 50)

    def test_split_by_words(self):
        out_dir = os.path.join(self.test_dir, "split_words")
        result = WordSplitter.split_document(
            self.sample_path,
            num_parts=3,
            output_dir=out_dir,
            method="words"
        )
        self.assertEqual(result["actual_parts"], 3)
        self.assertEqual(len(result["parts"]), 3)

        total_part_words = 0
        for p in result["parts"]:
            self.assertTrue(os.path.isfile(p["file_path"]))
            # Verify it opens cleanly with python-docx
            d = docx.Document(p["file_path"])
            total_part_words += p["words"]
            self.assertGreater(len(d.paragraphs) + len(d.tables), 0)

        # Word count of all parts should equal total source words
        self.assertEqual(total_part_words, result["total_source_words"])

    def test_split_by_blocks(self):
        out_dir = os.path.join(self.test_dir, "split_blocks")
        result = WordSplitter.split_document(
            self.sample_path,
            num_parts=2,
            output_dir=out_dir,
            method="blocks"
        )
        self.assertEqual(result["actual_parts"], 2)
        total_blocks = sum(p["block_count"] for p in result["parts"])
        self.assertEqual(total_blocks, result["total_source_blocks"])

    def test_split_by_headings(self):
        out_dir = os.path.join(self.test_dir, "split_headings")
        result = WordSplitter.split_document(
            self.sample_path,
            num_parts=3,
            output_dir=out_dir,
            method="headings"
        )
        self.assertEqual(result["actual_parts"], 3)
        for p in result["parts"]:
            self.assertTrue(os.path.isfile(p["file_path"]))

    def test_split_by_memory(self):
        out_dir = os.path.join(self.test_dir, "split_memory")
        result = WordSplitter.split_document(
            self.sample_path,
            num_parts=3,
            output_dir=out_dir,
            method="memory"
        )
        self.assertEqual(result["actual_parts"], 3)
        for p in result["parts"]:
            self.assertTrue(os.path.isfile(p["file_path"]))
            d = docx.Document(p["file_path"])
            self.assertGreater(len(d.paragraphs), 0)


if __name__ == "__main__":
    unittest.main()
